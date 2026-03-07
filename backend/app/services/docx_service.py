from pathlib import Path
from typing import Dict

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt

from ..database import STORAGE_DIR


def save_base_resume(file_bytes: bytes, filename: str) -> Path:
    resumes_dir = STORAGE_DIR / "resumes"
    resumes_dir.mkdir(parents=True, exist_ok=True)
    path = resumes_dir / filename
    path.write_bytes(file_bytes)
    return path


def extract_text_from_docx(path: Path) -> str:
    document = Document(str(path))
    parts = [para.text for para in document.paragraphs if para.text.strip()]
    return "\n".join(parts)


def _apply_classic_template(doc: Document) -> None:
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)


def _apply_modern_template(doc: Document) -> None:
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Segoe UI"
    font.size = Pt(10)
    if "Modern Heading" not in doc.styles:
        heading_style = doc.styles.add_style("Modern Heading", WD_STYLE_TYPE.PARAGRAPH)
        h_font = heading_style.font
        h_font.name = "Segoe UI Semibold"
        h_font.size = Pt(13)


def _apply_compact_template(doc: Document) -> None:
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Arial"
    font.size = Pt(10)


def generate_aligned_resume_docx(
    sections: Dict[str, str],
    filename: str = "aligned_resume.docx",
    template: str = "classic",
) -> Path:
    generated_dir = STORAGE_DIR / "generated"
    generated_dir.mkdir(parents=True, exist_ok=True)

    doc = Document()

    template_key = (template or "classic").lower()
    if template_key == "modern":
        _apply_modern_template(doc)
        heading_style_name = "Modern Heading"
    elif template_key == "compact":
        _apply_compact_template(doc)
        heading_style_name = "Heading 2"
    else:
        _apply_classic_template(doc)
        heading_style_name = "Heading 1"

    for section, content in sections.items():
        heading = doc.add_paragraph(section)
        heading.style = heading_style_name
        for line in content.split("\n"):
            if line.strip():
                para = doc.add_paragraph(line)
                para.style = "Normal"

    out_path = generated_dir / filename
    doc.save(str(out_path))
    return out_path

